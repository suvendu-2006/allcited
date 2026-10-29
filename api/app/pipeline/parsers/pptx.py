import io
from typing import List, Union

from pptx import Presentation

from app.pipeline.parsers.base import ParsedBlock


class PPTXParserError(Exception):
    """Raised when PPTX document cannot be read or parsed."""

    pass


def parse_pptx(
    content: Union[str, bytes],
    source_name: str = "presentation.pptx",
) -> List[ParsedBlock]:
    """
    Parse a PowerPoint (.pptx) file into structured ParsedBlocks.
    Each slide produces block(s) with slide title as section path,
    shape text/tables, and attached speaker notes.
    """
    if isinstance(content, str):
        with open(content, "rb") as f:
            bytes_data = f.read()
    else:
        bytes_data = content

    try:
        prs = Presentation(io.BytesIO(bytes_data))
    except Exception as e:
        raise PPTXParserError(f"Corrupt or unreadable PPTX presentation: {str(e)}") from e

    blocks: List[ParsedBlock] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        # 1. Extract slide title if present
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text.strip()
        else:
            slide_title = f"Slide {slide_idx}"

        section_path = [slide_title]
        slide_content_parts: List[str] = []

        # 2. Extract content from shapes
        for shape in slide.shapes:
            # Don't duplicate the title shape if it's already captured
            if shape == slide.shapes.title:
                continue

            # Standard text frames
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_content_parts.append(text)

            # Tables in slides
            elif shape.has_table:
                table = shape.table
                table_lines = []
                for r_idx, row in enumerate(table.rows):
                    cells = [
                        c.text.strip().replace("\n", " ").replace("|", "\\|") for c in row.cells
                    ]
                    if any(cells):
                        table_lines.append("| " + " | ".join(cells) + " |")
                        if r_idx == 0:
                            table_lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
                if table_lines:
                    slide_content_parts.append("\n".join(table_lines))

        # 3. Extract speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                slide_content_parts.append(f"[Speaker Notes: {notes_text}]")

        # 4. Construct block for this slide
        full_slide_text = "\n\n".join(slide_content_parts).strip()
        if not full_slide_text:
            # Even if slide has only a title, include the title as content
            full_slide_text = slide_title

        blocks.append(
            ParsedBlock(
                text=full_slide_text,
                section_path=section_path,
                page=slide_idx,
                source=source_name,
            )
        )

    return blocks
